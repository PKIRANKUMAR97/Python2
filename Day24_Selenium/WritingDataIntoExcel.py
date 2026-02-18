import openpyxl

file2 = r"E:\Pavan_Trainings\test_data.xlsx"
wb2=openpyxl.load_workbook(file2)
sh2=wb2['Sheet1']

# for r in range(1,6):
#     for c in range(1,6):
#         sh2.cell(r,c).value="Hello World"

sh2.cell(1,1).value="Kiran1"
sh2.cell(1,2).value="Kiran2"
sh2.cell(1,3).value="Kiran3"

sh2.cell(2,1).value="Kiran4"
sh2.cell(2,2).value="Kiran5"
sh2.cell(2,3).value="Kiran6"
wb2.save(file2)