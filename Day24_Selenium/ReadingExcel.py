import openpyxl
file = r"E:\Pavan_Trainings\test_data.xlsx"
wb=openpyxl.load_workbook(file)
sh=wb['Sheet']

rows=sh.max_row
cols=sh.max_column
for row in range(1,rows+1):
    for col in range(1,cols+1):
        print(sh.cell(row,col).value,end="      ")
    print()
